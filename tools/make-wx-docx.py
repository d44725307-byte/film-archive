#!/usr/bin/env python3
# 把文章 JSON 生成「可直接导入公众号」的 .docx（标题 + 正文 + 图片 + 图注，排版就位）
# 用法：python3 tools/make-wx-docx.py <文章.json> [<文章.json> ...] [--out 输出目录]
# 说明：公众号后台「图文素材 → 新建 → 文档导入」选 Word 即可，图片会一起进去。
import json, os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FONT = 'PingFang SC'          # 公众号常见字体；Word 里没有会自动回退
GOLD = RGBColor(0xA8, 0x7A, 0x10)
DARK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x99, 0x99, 0x99)
IMG_W = Cm(14.5)              # 正文图宽（A4 内边距内基本满宽）

def set_font(run, size=11, bold=False, color=DARK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {}); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT); rFonts.set(qn('w:hAnsi'), FONT); rFonts.set(qn('w:eastAsia'), FONT)

def para(doc, align=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.6
    return p

TOKEN = re.compile(r'(<strong>.*?</strong>|<a\s+href="[^"]*"[^>]*>.*?</a>|<br\s*/?>)', re.S)

def md2html(s):
    """Markdown 加粗/链接 → HTML（文章 JSON 里两种写法都有）"""
    s = re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', s or '')
    s = re.sub(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', r'<a href="\2">\1</a>', s)
    return s

def add_inline(p, html, size=11, color=DARK):
    """把简易 HTML（strong/a/br）转成 docx 文本片段"""
    for part in TOKEN.split(md2html(html)):
        if not part: continue
        if part.startswith('<br'):
            p.add_run().add_break()
        elif part.startswith('<strong>'):
            r = p.add_run(re.sub(r'<[^>]+>', '', part)); set_font(r, size, bold=True, color=color)
        elif part.startswith('<a '):
            m = re.match(r'<a\s+href="([^"]*)"[^>]*>(.*?)</a>', part, re.S)
            if m:
                txt = re.sub(r'<[^>]+>', '', m.group(2))
                r = p.add_run(txt); set_font(r, size, bold=True, color=GOLD)
                r2 = p.add_run('（' + m.group(1) + '）'); set_font(r2, size - 1.5, color=GREY)
        else:
            r = p.add_run(re.sub(r'<[^>]+>', '', part)); set_font(r, size, color=color)

def build(art, outdir, imgdir):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2); s.left_margin = s.right_margin = Cm(2.2)

    # 标题
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    set_font(p.add_run(art['title']), 17, bold=True, color=RGBColor(0x22, 0x22, 0x22))

    for b in art.get('blocks', []):
        t = b.get('t')
        if t == 'h2':
            p = para(doc, space_before=14, space_after=8)
            set_font(p.add_run(re.sub(r'<[^>]+>', '', b.get('html', ''))), 13.5, bold=True, color=GOLD)
        elif t in ('lead', 'p'):
            p = para(doc)
            add_inline(p, b.get('html', ''), 11.5 if t == 'lead' else 11)
        elif t == 'quote':
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=12)
            set_font(p.add_run(re.sub(r'<[^>]+>', '', b.get('html', ''))), 13, bold=True, color=RGBColor(0x22, 0x40, 0x2F))
        elif t == 'fig':
            src = os.path.join(imgdir, os.path.basename(b.get('src', '')))
            if os.path.exists(src):
                doc.add_picture(src, width=IMG_W)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b.get('caption'):
                    c = para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
                    set_font(c.add_run(b['caption']), 9, color=GREY)
            else:
                print('   ⚠️ 缺图:', src)
    out = os.path.join(outdir, art['slug'] + '.docx')
    doc.save(out)
    return out

def main():
    argv = sys.argv[1:]
    outdir = 'wx-docx'
    if '--out' in argv:
        i = argv.index('--out')
        outdir = argv[i + 1]
        del argv[i:i + 2]
    args = [a for a in argv if not a.startswith('--')]
    os.makedirs(outdir, exist_ok=True)
    for f in args:
        art = json.load(open(f, encoding='utf-8'))
        # 从文章 JSON 向上找含 samples/photos 的目录
        d = os.path.dirname(os.path.abspath(f))
        while d != '/' and not os.path.isdir(os.path.join(d, 'samples', 'photos')):
            d = os.path.dirname(d)
        imgdir = os.path.join(d, 'samples', 'photos')
        out = build(art, outdir, imgdir)
        print('✅', art['title'][:30], '→', out, str(round(os.path.getsize(out) / 1024)) + 'KB')

if __name__ == '__main__':
    main()
